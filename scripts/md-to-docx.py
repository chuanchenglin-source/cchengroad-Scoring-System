"""
Markdown → Word (.docx) 轉換工具
專為繁體中文文件設計，保留表格、標題層級、程式碼區塊
"""
import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding='utf-8')

# 字型設定
FONT_TC = '微軟正黑體'      # 繁體中文主字型
FONT_EN = 'Calibri'          # 英文字型
FONT_CODE = 'Consolas'       # 程式碼字型

def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)

def set_run_font(run, font_name=FONT_TC, font_size=11, bold=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    # 東亞字型設定（繁體中文需要）
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT_TC)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)

def add_header(doc, text):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, FONT_TC, 10, bold=True, color=RGBColor(0x44, 0x44, 0x44))

def add_heading(doc, text, level):
    sizes = {1: 20, 2: 16, 3: 14, 4: 12}
    p = doc.add_paragraph()
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, FONT_TC, sizes.get(level, 11), bold=True,
                 color=RGBColor(0x1E, 0x3A, 0x8A) if level <= 2 else RGBColor(0x37, 0x41, 0x51))
    p.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
    p.paragraph_format.space_after = Pt(6)

def parse_inline(text):
    """回傳 [(text, bold, code), ...] 以支援 **粗體**、`程式碼`"""
    parts = []
    i = 0
    while i < len(text):
        # 粗體 **xxx**
        m = re.match(r'\*\*(.+?)\*\*', text[i:])
        if m:
            parts.append((m.group(1), True, False))
            i += m.end()
            continue
        # 行內程式碼 `xxx`
        m = re.match(r'`([^`]+)`', text[i:])
        if m:
            parts.append((m.group(1), False, True))
            i += m.end()
            continue
        # 一般字元
        parts.append((text[i], False, False))
        i += 1
    # 合併相鄰相同格式
    merged = []
    for txt, bold, code in parts:
        if merged and merged[-1][1] == bold and merged[-1][2] == code:
            merged[-1] = (merged[-1][0] + txt, bold, code)
        else:
            merged.append((txt, bold, code))
    return merged

def add_paragraph_inline(doc, text, list_type=None):
    p = doc.add_paragraph()
    if list_type == 'bullet':
        p.style = 'List Bullet'
    elif list_type == 'number':
        p.style = 'List Number'
    p.paragraph_format.space_after = Pt(4)
    for txt, bold, code in parse_inline(text):
        run = p.add_run(txt)
        if code:
            set_run_font(run, FONT_CODE, 10, bold=bold, color=RGBColor(0xBE, 0x18, 0x5D))
        else:
            set_run_font(run, FONT_TC, 11, bold=bold)

def add_code_block(doc, code_lines, lang=''):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    # 灰底
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F3F4F6')
    pPr.append(shd)
    run = p.add_run('\n'.join(code_lines))
    set_run_font(run, FONT_CODE, 9, color=RGBColor(0x37, 0x41, 0x51))

def add_table(doc, rows):
    """rows 是 list[list[str]]，第 0 列視為 header"""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    for r in rows:
        while len(r) < ncols:
            r.append('')
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Light Grid Accent 1'
    table.autofit = True

    for i, row in enumerate(rows):
        tr = table.rows[i]
        for j, cell_text in enumerate(row):
            cell = tr.cells[j]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            # 清空預設段落
            cell.paragraphs[0].clear()
            p = cell.paragraphs[0]
            for txt, bold, code in parse_inline(cell_text):
                run = p.add_run(txt)
                if code:
                    set_run_font(run, FONT_CODE, 9, bold=bold or (i == 0),
                                 color=RGBColor(0xBE, 0x18, 0x5D))
                else:
                    set_run_font(run, FONT_TC, 10, bold=bold or (i == 0))
            if i == 0:
                set_cell_shading(cell, 'DBEAFE')  # 標題列淺藍底
    # 加一點段後空白
    doc.add_paragraph()

def convert(md_path: Path, docx_path: Path):
    text = md_path.read_text(encoding='utf-8')
    lines = text.split('\n')

    doc = Document()
    # 預設字型
    style = doc.styles['Normal']
    style.font.name = FONT_TC
    style.font.size = Pt(11)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT_TC)
    rFonts.set(qn('w:ascii'), FONT_TC)
    rFonts.set(qn('w:hAnsi'), FONT_TC)

    # 頁邊距
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # 頁首
    add_header(doc, '海洋親證班計分系統 — Supabase 遷移進度匯報')

    i = 0
    while i < len(lines):
        line = lines[i]

        # 程式碼區塊
        if line.startswith('```'):
            lang = line[3:].strip()
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i])
                i += 1
            add_code_block(doc, code, lang)
            i += 1
            continue

        # 標題
        m = re.match(r'^(#{1,4})\s+(.+)$', line)
        if m:
            add_heading(doc, m.group(2).strip(), len(m.group(1)))
            i += 1
            continue

        # 水平線
        if re.match(r'^---+$', line.strip()):
            p = doc.add_paragraph()
            pPr = p._element.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:color'), 'CCCCCC')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # 表格
        if '|' in line and i + 1 < len(lines) and re.match(r'^\s*\|?[\s\-:\|]+\|?\s*$', lines[i + 1]):
            rows = []
            while i < len(lines) and '|' in lines[i]:
                if re.match(r'^\s*\|?[\s\-:\|]+\|?\s*$', lines[i]):
                    i += 1
                    continue
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        # 項目清單（bullet）
        m = re.match(r'^(\s*)[-*]\s+(.+)$', line)
        if m:
            add_paragraph_inline(doc, m.group(2), 'bullet')
            i += 1
            continue

        # 編號清單
        m = re.match(r'^\s*\d+\.\s+(.+)$', line)
        if m:
            add_paragraph_inline(doc, m.group(1), 'number')
            i += 1
            continue

        # 空行
        if line.strip() == '':
            i += 1
            continue

        # 引用
        if line.startswith('>'):
            content = line.lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run(content)
            set_run_font(run, FONT_TC, 10, color=RGBColor(0x6B, 0x72, 0x80))
            i += 1
            continue

        # 一般段落
        add_paragraph_inline(doc, line)
        i += 1

    doc.save(str(docx_path))
    print(f'✅ 產出：{docx_path}')

if __name__ == '__main__':
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else None
    if not src or not dst:
        print('用法：python md-to-docx.py <input.md> <output.docx>')
        sys.exit(1)
    convert(src, dst)
